from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth import require_report_access
from app.database import get_db
from app.models import Checkout, InventoryItem, ItemCategory, Movement, User
from app.services.inventory import availability_label, available_qty

router = APIRouter(tags=["reports"])


@router.get("/reports/summary")
def report_summary(db: Session = Depends(get_db), _user: User = Depends(require_report_access)):
    items = db.execute(
        select(InventoryItem)
        .where(InventoryItem.category != ItemCategory.sops)
        .order_by(InventoryItem.name)
    ).scalars().all()
    item_ids = {item.id for item in items}
    open_tx = (
        db.execute(
            select(Checkout).where(
                Checkout.returned_at.is_(None), Checkout.item_id.in_(item_ids)
            )
        ).scalars().all()
    )
    overdue = []
    now = datetime.now(timezone.utc)
    for c in open_tx:
        if c.expected_return and c.expected_return < now:
            item = db.get(InventoryItem, c.item_id)
            overdue.append(
                {
                    "txId": c.tx_id,
                    "item": item.name if item else "?",
                    "serial_number": item.serial_number if item else None,
                    "personRole": c.person_role,
                    "expectedReturn": c.expected_return.isoformat(),
                }
            )
    low = [
        {
            "item": it.name,
            "serial_number": it.serial_number,
            "quantity": it.qty_on_hand,
            "available": available_qty(db, it),
            "reorder_min": it.reorder_min,
        }
        for it in items
        if available_qty(db, it) <= it.reorder_min
    ]
    return {
        "ok": True,
        "open_count": len(open_tx),
        "overdue": overdue,
        "low_stock": low,
    }


@router.get("/exports/inventory.xlsx")
def export_inventory_xlsx(db: Session = Depends(get_db), _user: User = Depends(require_report_access)):
    wb = Workbook()

    ws = wb.active
    ws.title = "Stock"
    ws.append(["SKU", "Article", "N° de série", "Catégorie", "Quantité", "Disponible", "Seuil de commande", "Sous le seuil", "Code-barres"])
    items = db.execute(
        select(InventoryItem)
        .where(InventoryItem.category != ItemCategory.sops)
        .order_by(InventoryItem.name)
    ).scalars().all()
    item_ids = {item.id for item in items}
    for it in items:
        avail = available_qty(db, it)
        ws.append(
            [
                it.sku,
                it.name,
                it.serial_number or "",
                it.category.value,
                it.qty_on_hand,
                avail,
                it.reorder_min,
                "OUI" if avail <= it.reorder_min else "NON",
                it.barcode or "",
            ]
        )

    ws2 = wb.create_sheet("Outils empruntés")
    ws2.append(["Transaction", "Article", "N° de série", "Pris par", "Date de sortie", "Retour prévu"])
    opens = db.execute(
        select(Checkout).where(
            Checkout.returned_at.is_(None), Checkout.item_id.in_(item_ids)
        )
    ).scalars().all()
    for c in opens:
        item = db.get(InventoryItem, c.item_id)
        ws2.append(
            [
                c.tx_id,
                item.name if item else "",
                item.serial_number if item and item.serial_number else "",
                c.person_role,
                c.taken_at.isoformat() if c.taken_at else "",
                c.expected_return.isoformat() if c.expected_return else "",
            ]
        )

    ws3 = wb.create_sheet("Mouvements récents")
    ws3.append(["Date", "Type", "Article", "N° de série", "Quantité", "Utilisateur", "Motif", "Transaction"])
    moves = (
        db.execute(select(Movement).order_by(Movement.created_at.desc()).limit(500)).scalars().all()
    )
    for m in moves:
        item = db.get(InventoryItem, m.item_id) if m.item_id else None
        ws3.append(
            [
                m.created_at.isoformat() if m.created_at else "",
                m.movement_type.value,
                m.item_name,
                item.serial_number if item and item.serial_number else "",
                m.qty,
                m.actor,
                m.reason or "",
                m.related_tx_id or "",
            ]
        )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"stock-kia-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
